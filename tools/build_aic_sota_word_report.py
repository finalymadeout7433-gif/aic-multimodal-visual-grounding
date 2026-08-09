from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\12525\Documents\pytorch\baseline_v0")
OUT = ROOT / "reports" / "AIC_近年多模态语义目标定位论文与开源模型调研_2026-08-01.docx"

NAVY = "17365D"
BLUE = "2E74B5"
BLUE_DARK = "1F4D78"
TEAL = "2A7F9E"
LIGHT_BLUE = "EAF2F8"
LIGHT_TEAL = "E7F3F5"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "D9E1E8"
DARK_GRAY = "3F4A56"
GREEN = "2E7D32"
LIGHT_GREEN = "EAF4EA"
AMBER = "A15C00"
LIGHT_AMBER = "FFF3D6"
RED = "A62C2B"
LIGHT_RED = "FCE8E6"
WHITE = "FFFFFF"
BLACK = "1F2933"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        # Keep OOXML child order valid. Appending shading after tcMar/vAlign
        # makes Word open the file in repair mode on documents with many tables.
        tc_pr.insert_element_before(
            shd,
            "w:noWrap",
            "w:tcMar",
            "w:textDirection",
            "w:tcFitText",
            "w:vAlign",
            "w:hideMark",
            "w:headers",
        )
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.insert_element_before(
            tc_mar,
            "w:textDirection",
            "w:tcFitText",
            "w:vAlign",
            "w:hideMark",
            "w:headers",
        )
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, **edges) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge_name, edge_data in edges.items():
        tag = f"w:{edge_name}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        for key, value in edge_data.items():
            element.set(qn(f"w:{key}"), str(value))


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_width(table, width_twips: int = 9360) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_twips))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.insert_element_before(
            tbl_ind,
            "w:tblBorders",
            "w:shd",
            "w:tblLayout",
            "w:tblCellMar",
            "w:tblLook",
            "w:tblCaption",
            "w:tblDescription",
        )
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")


def set_east_asia_font(run, name: str = "Microsoft YaHei") -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def add_hyperlink(paragraph, text: str, url: str, color: str = BLUE, underline: bool = True):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        r_pr.append(u)
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Calibri")
    r_fonts.set(qn("w:hAnsi"), "Calibri")
    r_fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r_pr.append(r_fonts)
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instruction
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in (
        ("Title", 25, NAVY, 0, 14),
        ("Subtitle", 12, DARK_GRAY, 0, 8),
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, BLUE_DARK, 8, 4),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = style_name != "Subtitle"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Table Text" not in styles:
        s = styles.add_style("Table Text", WD_STYLE_TYPE.PARAGRAPH)
        s.base_style = styles["Normal"]
        s.font.name = "Calibri"
        s.font.size = Pt(9)
        s._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        s.paragraph_format.space_after = Pt(2)
        s.paragraph_format.line_spacing = 1.05

    if "Small Note" not in styles:
        s = styles.add_style("Small Note", WD_STYLE_TYPE.PARAGRAPH)
        s.base_style = styles["Normal"]
        s.font.name = "Calibri"
        s.font.size = Pt(9)
        s.font.color.rgb = RGBColor.from_string(DARK_GRAY)
        s._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        s.paragraph_format.space_after = Pt(4)
        s.paragraph_format.line_spacing = 1.05

    if "Code Inline" not in styles:
        s = styles.add_style("Code Inline", WD_STYLE_TYPE.CHARACTER)
        s.font.name = "Consolas"
        s.font.size = Pt(9)
        s.font.color.rgb = RGBColor.from_string(BLUE_DARK)
        s._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def configure_page(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def add_header_footer(section) -> None:
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("AIC 多模态语义目标定位｜论文与开源模型调研")
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string(DARK_GRAY)
    set_east_asia_font(run)
    p.paragraph_format.space_after = Pt(0)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("2026-08-01   ·   ")
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string(DARK_GRAY)
    set_east_asia_font(run)
    add_field(p, " PAGE ")
    run = p.add_run(" / ")
    set_east_asia_font(run)
    add_field(p, " NUMPAGES ")


def add_cover(doc: Document) -> None:
    band = doc.add_table(rows=1, cols=1)
    band.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(band)
    cell = band.cell(0, 0)
    set_cell_shading(cell, NAVY)
    set_cell_margins(cell, top=80, bottom=80)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("RESEARCH & IMPLEMENTATION BRIEF")
    r.font.size = Pt(9)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(WHITE)
    set_east_asia_font(r)

    for _ in range(4):
        doc.add_paragraph("")

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("AIC 赛题一")
    set_east_asia_font(r)
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("近年多模态语义目标定位论文与开源模型调研")
    set_east_asia_font(r)

    rule = doc.add_table(rows=1, cols=1)
    rule.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(rule)
    c = rule.cell(0, 0)
    set_cell_shading(c, TEAL)
    set_cell_margins(c, top=18, bottom=18)
    c.paragraphs[0].paragraph_format.space_after = Pt(0)

    p = doc.add_paragraph(style="Subtitle")
    p.paragraph_format.space_before = Pt(14)
    p.add_run("面向 RGB + Infrared + Depth + English Query → Normalized BBox 的技术选型与落地路线")
    p = doc.add_paragraph(style="Subtitle")
    p.add_run("基于官方论文、官方代码仓库、模型卡与本项目真实平台结果核验")

    doc.add_paragraph("")
    meta = doc.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_width(meta)
    labels = ["项目目录", "证据截止", "目标硬件", "报告用途"]
    values = [
        r"D:\12525\Documents\pytorch\baseline_v0",
        "2026-08-01（Asia/Shanghai）",
        "RTX 4060 Laptop GPU（优先考虑 8 GB 级显存约束）",
        "供下一轮模型验证、训练路线选择及 ChatGPT 网页端继续分析",
    ]
    for i, (label, value) in enumerate(zip(labels, values)):
        meta.cell(i, 0).width = Inches(1.35)
        meta.cell(i, 1).width = Inches(5.15)
        set_cell_shading(meta.cell(i, 0), LIGHT_BLUE)
        set_cell_margins(meta.cell(i, 0))
        set_cell_margins(meta.cell(i, 1))
        p0 = meta.cell(i, 0).paragraphs[0]
        p0.style = "Table Text"
        r0 = p0.add_run(label)
        r0.bold = True
        r0.font.color.rgb = RGBColor.from_string(BLUE_DARK)
        set_east_asia_font(r0)
        p1 = meta.cell(i, 1).paragraphs[0]
        p1.style = "Table Text"
        r1 = p1.add_run(value)
        set_east_asia_font(r1)
    doc.add_paragraph("")
    p = doc.add_paragraph(style="Small Note")
    p.add_run("说明：本文把“论文正式发表状态”“代码/数据/权重是否可获取”“能否在当前硬件上复现”分开陈述，不把预印本或仓库自述等同于已验证可落地。")
    doc.add_page_break()


def add_callout(doc: Document, title: str, body: str, color=BLUE, fill=LIGHT_BLUE) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(table)
    table.columns[0].width = Inches(0.12)
    table.columns[1].width = Inches(6.38)
    left = table.cell(0, 0)
    right = table.cell(0, 1)
    set_cell_shading(left, color)
    set_cell_shading(right, fill)
    set_cell_margins(left, top=70, bottom=70, start=20, end=20)
    set_cell_margins(right, top=110, bottom=100, start=150, end=150)
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = right.paragraphs[0]
    p.style = "Table Text"
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor.from_string(color)
    set_east_asia_font(r)
    p2 = right.add_paragraph(style="Table Text")
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(body)
    r2.font.size = Pt(10)
    set_east_asia_font(r2)


def add_bullet(doc: Document, text: str, level: int = 0, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.left_indent = Inches(0.25 + 0.25 * level)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.10
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        set_east_asia_font(r)
        r2 = p.add_run(text[len(bold_prefix):])
        set_east_asia_font(r2)
    else:
        r = p.add_run(text)
        set_east_asia_font(r)


def add_numbered(doc: Document, number: int, text: str) -> None:
    """Add an explicitly numbered item so each section starts from 1."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(5)
    prefix = p.add_run(f"{number}. ")
    prefix.bold = True
    prefix.font.color.rgb = RGBColor.from_string(BLUE_DARK)
    set_east_asia_font(prefix)
    r = p.add_run(text)
    set_east_asia_font(r)


def add_body(doc: Document, text: str, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        set_east_asia_font(r)
        r2 = p.add_run(text[len(bold_prefix):])
        set_east_asia_font(r2)
    else:
        r = p.add_run(text)
        set_east_asia_font(r)


def add_table(doc: Document, headers: Sequence[str], rows: Iterable[Sequence[str]], widths: Sequence[float] | None = None, font_size=8.5):
    rows = list(rows)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_width(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for j, header in enumerate(headers):
        cell = hdr.cells[j]
        if widths:
            cell.width = Inches(widths[j])
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.style = "Table Text"
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(header)
        r.bold = True
        r.font.size = Pt(font_size)
        r.font.color.rgb = RGBColor.from_string(WHITE)
        set_east_asia_font(r)
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j, text in enumerate(row):
            if widths:
                cells[j].width = Inches(widths[j])
            set_cell_margins(cells[j])
            cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if i % 2 == 1:
                set_cell_shading(cells[j], "FAFBFC")
            p = cells[j].paragraphs[0]
            p.style = "Table Text"
            r = p.add_run(str(text))
            r.font.size = Pt(font_size)
            set_east_asia_font(r)
    return table


def add_key_value_table(doc: Document, rows: Sequence[tuple[str, str]]) -> None:
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_width(table)
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].width = Inches(1.45)
        cells[1].width = Inches(5.05)
        set_cell_shading(cells[0], LIGHT_GRAY)
        for c in cells:
            set_cell_margins(c)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        p0 = cells[0].paragraphs[0]
        p0.style = "Table Text"
        r0 = p0.add_run(label)
        r0.bold = True
        r0.font.color.rgb = RGBColor.from_string(BLUE_DARK)
        set_east_asia_font(r0)
        p1 = cells[1].paragraphs[0]
        p1.style = "Table Text"
        r1 = p1.add_run(value)
        set_east_asia_font(r1)


def add_link_line(doc: Document, links: Sequence[tuple[str, str]]) -> None:
    p = doc.add_paragraph(style="Small Note")
    p.paragraph_format.space_before = Pt(2)
    r = p.add_run("官方链接：")
    r.bold = True
    set_east_asia_font(r)
    for idx, (label, url) in enumerate(links):
        if idx:
            r2 = p.add_run("  ·  ")
            set_east_asia_font(r2)
        add_hyperlink(p, label, url)


def add_paper_section(doc: Document, number: str, title: str, facts: Sequence[tuple[str, str]], summary: str,
                      relevance: Sequence[str], integration: Sequence[str], risks: Sequence[str],
                      links: Sequence[tuple[str, str]], verdict: str, verdict_color=TEAL, verdict_fill=LIGHT_TEAL) -> None:
    doc.add_heading(f"{number}  {title}", level=2)
    add_key_value_table(doc, facts)
    add_body(doc, summary)
    doc.add_heading("与 AIC 的相关性", level=3)
    for item in relevance:
        add_bullet(doc, item)
    doc.add_heading("推荐接入方式", level=3)
    for idx, item in enumerate(integration, start=1):
        add_numbered(doc, idx, item)
    doc.add_heading("主要风险与限制", level=3)
    for item in risks:
        add_bullet(doc, item)
    add_callout(doc, "结论", verdict, color=verdict_color, fill=verdict_fill)
    add_link_line(doc, links)


def build_document() -> Document:
    doc = Document()
    configure_styles(doc)
    for section in doc.sections:
        configure_page(section)
        add_header_footer(section)
    add_cover(doc)

    doc.add_heading("执行摘要", level=1)
    add_callout(
        doc,
        "一句话结论",
        "下一轮不应继续扩大当前 RefCOCO-LightGBM Ranker；最值得优先验证的是 PIZA（极小目标）、APE-Ti（结构/区域目标）、MM-Grounding-DINO-Tiny（可训练基座）和 RGBT-VGNet（RGB+热红外长期主线）。YOLO-World、OmDet-Turbo、OWLv2、OV-DINO更适合作为补充候选源，而不是单独承担复杂指代表达推理。",
        color=GREEN,
        fill=LIGHT_GREEN,
    )
    add_body(doc, "本报告对网页端列出的论文和模型重新进行了官方来源核验，并把“论文水平、任务是否真正匹配、数据/代码/权重是否公开、对 RTX 4060 的可落地性”分开判断。以下结论以 2026-08-01 可访问的论文主页、官方仓库和模型卡为准。")

    doc.add_heading("当前项目的真实证据边界", level=2)
    add_table(
        doc,
        ["提交", "系统", "AIC ACC@0.5", "结论"],
        [
            ("S01", "Florence-2 RGB-only", "0.4980", "当前稳定基线"),
            ("S02", "GroundingDINO-Tiny 原始 Top-1", "0.4938", "与 S01 接近，可保留为候选源"),
            ("S03", "GroundingDINO + RefCOCO Spatial LTR", "0.3190", "严重负迁移，不能继续作为默认提交"),
        ],
        widths=[0.65, 2.65, 1.15, 2.05],
        font_size=9,
    )
    add_bullet(doc, "本地 RefCOCO holdout 上的 0.6333 和 Top-10 oracle 0.8936 只代表外部同分布结果，不代表 AIC 的候选上限。")
    add_bullet(doc, "S03 平台下降已排除坐标、索引和 ZIP 装包错误，主要是排序器选择行为真实失效：偏向大框，并把目标替换成参照物。")
    add_bullet(doc, "因此本轮选型的核心不是再寻找一个在 RefCOCO 上更高的分数，而是补齐 AIC 的小目标、长关系句、结构区域、低光与跨模态分布。")

    doc.add_heading("优先级结论", level=2)
    add_table(
        doc,
        ["优先级", "项目", "建议角色", "为什么"],
        [
            ("P0", "PIZA + SOREC", "小目标专用分支", "直接针对极小目标 REC；代码、标注和 adapter 权重公开"),
            ("P0", "APE-Ti", "区域/结构目标对照", "支持 foreground、background stuff 与自然语言描述；轻量 checkpoint 已公开"),
            ("P0", "MM-GDINO-T", "下一代 RGB+文本可训练基座", "训练链路完整，Tiny 权重公开，便于受控微调"),
            ("P1", "RGBT-GroundBench / VGNet", "RGB+IR 主线", "目前与 AIC 的 RGB+热红外+Query 最接近；数据、代码和 checkpoint 资源公开"),
            ("P1", "OmDet-Turbo / YOLO-World / OWLv2", "低成本候选补充", "部署较轻、候选互补；但关系推理能力有限"),
            ("P2", "LLMDet Swin-T", "长语义开放词汇对照", "语义监督强，但训练和工程成本高，且不是专门 REC"),
            ("P3", "RGBDT500 / RDTTrack", "Depth+IR 融合参考", "视觉模态最接近，但任务是 tracking，且必须先排查数据重叠"),
            ("暂缓", "Thermo-VL", "架构参考", "7B VLM、VQA 目标、无 bbox head；代码/权重链接尚不能完整核验"),
        ],
        widths=[0.65, 1.55, 1.65, 2.65],
        font_size=8.3,
    )

    doc.add_heading("目录", level=1)
    for item in (
        "1. 网页端原分析的核验与修正",
        "2. 核心论文与数据集逐项分析",
        "3. 其他开源候选模型对照",
        "4. 面向 AIC 的组合方案与训练路线",
        "5. 推荐实验顺序、验收标准与风险控制",
        "6. 官方来源与链接索引",
    ):
        add_bullet(doc, item)
    doc.add_page_break()

    doc.add_heading("1. 网页端原分析的核验与修正", level=1)
    add_table(
        doc,
        ["原判断", "核验结果", "在本报告中的处理"],
        [
            ("RGBT-VGNet 可能没有公开权重", "需要更新：论文明确称 checkpoints 已公开，且存在 Hugging Face 模型仓库；但仓库暂无完整 model card。", "标记为“权重资源已公开，使用前需核对文件—配置映射和哈希”。"),
            ("Thermo-VL 是近期权威模型且代码/数据公开", "只确认 2026 年 arXiv 预印本和项目页；项目页显示 Code/Data 字样，但当前未解析到可用仓库链接，模型权重也未确认。", "降级为架构参考，不列为可立即复现模型。"),
            ("RGBDT500 可用于 AIC 三模态训练", "数据模态接近，但原任务是单目标跟踪，没有语言 Query。", "仅作为融合预训练/架构参考；先做许可和测试集近重复审计。"),
            ("OmDet-Turbo 属于顶会模型", "官方仓库对应 2024 arXiv 预印本；不是本次核验中已确认的顶会正式论文。", "放入“工程型开源候选模型”，不与 CVPR/ICCV/NeurIPS 正式论文混写。"),
            ("RefCOCO Top-10 oracle 表明 AIC 仍有巨大排序空间", "平台结果已证明外部域的 oracle 与排序规律不能直接迁移到 AIC。", "所有本地与平台分数分栏呈现；不再把 0.8936 解释为 AIC 理论上限。"),
        ],
        widths=[1.65, 2.8, 2.05],
        font_size=8.2,
    )
    add_callout(doc, "证据等级说明", "“正式发表”只用于官方会议/论文页面可核验的项目；“权重公开”要求存在可访问模型仓库或官方 checkpoint；仅有论文中“将公开”或项目页文本，不等同于已经可下载。", color=AMBER, fill=LIGHT_AMBER)

    doc.add_heading("2. 核心论文、数据集与方法", level=1)

    add_paper_section(
        doc,
        "2.1",
        "RGBT-GroundBench 与 RGBT-VGNet",
        [
            ("发表状态", "ECCV 2026（官方仓库标注；arXiv v2 于 2026-06 更新）"),
            ("任务", "RGB / TIR / RGB+TIR + referring expression → target bbox"),
            ("核心目标", "在低光、恶劣天气、遮挡、小/远目标条件下建立跨光谱 Visual Grounding 基准，并提供可靠性感知融合基线。"),
            ("数据规模", "21,535 对对齐 RGB-TIR，38,760 条 grounding 实例；train 26,604、val 2,032、test 10,124。"),
            ("训练集", "公开；由 FLIR、M3FD、MFAD 构建，含 Query、bbox 和条件标签。"),
            ("代码", "公开；提供训练、评测和可视化脚本。"),
            ("权重", "公开资源存在：论文称 checkpoints 已发布，Hugging Face 有 RGBT-Ground-Model 仓库；模型卡不完整，需先核对文件和配置。"),
            ("许可", "仓库声明仅用于学术研究；源数据仍受各自许可约束，重新分发/商用前需逐项确认。"),
            ("落地难度", "中高（4/5）：双视觉分支、对齐输入、LoRA 与跨模态融合；比单模型推理复杂。"),
        ],
        "RGBT-VGNet 的三个关键设计是：对 TIR 分支使用更强的非对称模态适配；让 Query 引导 RGB/TIR 的信息交互；再用模态可靠性、光照与语言先验融合两路特征。这一抽象与 AIC 的“什么时候应该相信红外”高度一致。",
        [
            "高度相关：它是当前核验到最接近“RGB + 热红外 + 语言 → bbox”的正式研究。",
            "可直接覆盖 AIC 的低光、人/车热目标和小目标场景，但不能处理 Depth。",
            "它提供比 RefCOCO 更接近 AIC 的训练/验证分布，有望降低现有 Ranker 的目标域负迁移。",
        ],
        [
            "先下载 annotation、代码和模型仓库元数据，固定 commit、配置与 checkpoint 哈希，不直接全量训练。",
            "在官方 val 上复现 RGB-only、TIR-only、RGB+TIR 三个结果，确认坐标与输入对齐链路。",
            "把 AIC 的 Infrared 映射到 TIR 分支，仅做无标签推理；输出仍映射到 RGB 坐标。",
            "若复现稳定，再以 RGBT-GroundBench train 进行 LoRA/adapter 微调；同时保留 Florence/GDINO 作为 RGB 对照。",
            "最终把 Depth 作为独立候选级晚融合信号，而不是修改该模型的第一版结构。",
        ],
        [
            "源数据许可不是统一宽松许可证，必须审计 FLIR/M3FD/MFAD 的下载与再使用条款。",
            "必须对 AIC 正式图像做 SHA-256、pHash/视觉 embedding 近重复检查，发现疑似同源帧即停止训练用途。",
            "Hugging Face 模型仓库暂无完整 model card，不能在未核对文件对应关系前直接运行。",
        ],
        [
            ("论文", "https://arxiv.org/abs/2512.24561"),
            ("官方代码", "https://github.com/crazyxiaoxi/RGBT-GroundBench"),
            ("数据", "https://huggingface.co/datasets/JiawenXi/RGBT-Ground-Dataset"),
            ("模型仓库", "https://huggingface.co/JiawenXi/RGBT-Ground-Model"),
        ],
        "长期主线价值最高，但不是下一次平台提交的最快单变量。建议先做资产审计和官方 val 复现，再决定是否训练。",
    )

    add_paper_section(
        doc,
        "2.2",
        "PIZA 与 SOREC：面向极小目标的 Referring Expression Comprehension",
        [
            ("发表状态", "ICCV 2025 正式论文"),
            ("任务", "RGB + referring expression → small-object bbox"),
            ("核心目标", "解决目标仅占图像极小面积时，常规 Visual Grounding 在缩放后看不清目标的问题。"),
            ("数据规模", "SOREC 共 100,000 条表达-bbox；Train-S 10,000、Train-L 61,369、Val 10,712、Test-A 10,815、Test-B 17,104。"),
            ("训练集", "标注 JSON 公开；底图来自 SODA-D，需按 SODA-D 官方方式下载。"),
            ("代码", "公开，基于 MMDetection/MM-Grounding-DINO。"),
            ("权重", "piza_adapter.pth 已在官方仓库；底层 MM-GDINO checkpoint 也有公开下载。"),
            ("许可", "代码与模型权重 MIT；底图仍遵循 SODA-D 条款。"),
            ("落地难度", "中（2.5/5）：有现成 adapter 和 demo，先推理很容易；混合数据再训练属于中等难度。"),
        ],
        "PIZA 不是盲目 2×2 Tile，而是参数高效的渐进式、迭代式放大：先找到可能区域，再逐步缩小搜索范围。它直接针对上一轮 Tile 实验暴露出的“小目标召回有效但计算昂贵”问题。",
        [
            "非常相关：AIC 中的无人机、摄像头、灯泡、标志和远距离目标与 SOREC 的极小目标难点一致。",
            "PIZA 仍是 RGB-only，不能替代 IR/Depth；但可以作为高置信小目标门控分支。",
            "它的训练数据包含长关系句与道路场景，比单纯 RefCOCO 更接近 AIC 的语言长度和目标尺度。",
        ],
        [
            "先用官方 adapter 在 SOREC val 复现，并对 AIC 官方小灯泡 sanity sample 做流程检查。",
            "在现有 9,555 条 AIC 推理中，只对小目标实体词、低候选置信度或全图候选过大的 Query 触发 PIZA。",
            "保留原模型 Top-1；只有 PIZA 输出通过 score、目标角色和框面积保护时才切换。",
            "若单变量平台有效，再把 SOREC 与 RefCOCO/gRefCOCO 做平衡采样，训练新的 adapter，而不是全量解冻。",
        ],
        [
            "SOREC 是驾驶场景，仍可能与 AIC 建筑区域、室内结构存在域差异。",
            "渐进缩放会增加延迟，必须记录触发率和每个 zoom step 的收益。",
            "必须避免把参照物放大后当作最终目标，target/reference 角色保护仍然必要。",
        ],
        [
            ("ICCV 论文", "https://openaccess.thecvf.com/content/ICCV2025/papers/Goto_Referring_Expression_Comprehension_for_Small_Objects_ICCV_2025_paper.pdf"),
            ("官方代码/标注/权重", "https://github.com/mmaiLab/sorec"),
        ],
        "这是下一轮最值得先做的单变量实验：现成权重、目标明确、与 AIC 极小目标高度匹配，且不需要马上训练大模型。",
        verdict_color=GREEN,
        verdict_fill=LIGHT_GREEN,
    )

    add_paper_section(
        doc,
        "2.3",
        "APE：Aligning and Prompting Everything All at Once",
        [
            ("发表状态", "CVPR 2024 正式论文"),
            ("任务", "统一目标检测、实例/语义分割、Visual Grounding / REC"),
            ("核心目标", "用一个模型同时理解前景物体、背景 stuff、区域和自然语言描述。"),
            ("训练数据", "使用 COCO、LVIS、Objects365、OpenImages、Visual Genome、SA-1B、RefCOCO、GQA、PhraseCut、Flickr30k 等多源数据。"),
            ("训练集", "训练配置公开；底层数据需要分别按官方条款获取。"),
            ("代码", "训练、推理和 demo 均公开。"),
            ("权重", "APE-L_A/B/C/D 与 APE-Ti 均有 Hugging Face checkpoint。"),
            ("许可", "Apache-2.0；外部训练数据许可证另行适用。"),
            ("落地难度", "中（3/5）：APE-Ti 可先推理；Detectron2/多数据配置环境比 Transformers 单模型复杂。"),
        ],
        "APE 的重要价值不是简单“更强检测器”，而是把普通物体之外的 background stuff、结构区域和自然语言描述也纳入统一对齐。官方仓库称一个模型在 160 个数据集上达到有竞争力表现，并提供只有 6M backbone 的 APE-Ti。",
        [
            "高度适合 AIC 中的 passage、awning、entrance、road、promotional area、building structure 等非 COCO 标准类别。",
            "可作为 Florence/GDINO 之外的区域型候选生成器，尤其适合 candidate union。",
            "不含 IR/Depth，因此是 RGB 语义/区域分支，不是最终三模态模型。",
        ],
        [
            "先运行 APE-Ti，不从 APE-L 开始；使用原始 Query 与目标短语两种提示分别产生候选。",
            "在 RefCOCO、gRefCOCO 与区域/结构验证子集上比较 Top-1 和 candidate union oracle。",
            "对 AIC 只生成一份单模型 Top-1 对照提交，或将 APE 候选加入保守 union；一次提交只改变一个变量。",
            "若区域类显著互补，再考虑冻结 backbone，仅训练 grounding/fusion 相关模块。",
        ],
        [
            "APE 预训练数据非常广，外部验证可能存在重叠；必须记录 checkpoint 的训练数据。",
            "APE-Ti 的“6M backbone”不等于整个模型只有 6M 参数，不能据此承诺极低显存。",
            "全量训练配置面向多 GPU；RTX 4060 更适合推理、LoRA 或局部微调。",
        ],
        [
            ("CVPR 论文", "https://openaccess.thecvf.com/content/CVPR2024/html/Shen_Aligning_and_Prompting_Everything_All_at_Once_for_Universal_Visual_CVPR_2024_paper.html"),
            ("官方代码/权重", "https://github.com/shenyunhang/APE"),
        ],
        "P0 横向验证模型。它最可能补齐当前系统对区域、建筑结构和背景类目标的弱点。",
        verdict_color=GREEN,
        verdict_fill=LIGHT_GREEN,
    )

    add_paper_section(
        doc,
        "2.4",
        "LLMDet：利用大语言模型监督开放词汇检测器",
        [
            ("发表状态", "CVPR 2025 Highlight"),
            ("任务", "RGB + text vocabulary/description → open-vocabulary detection bbox"),
            ("核心目标", "将 grounding loss 与图像级/区域级描述生成联合训练，让检测器吸收更丰富的语义监督。"),
            ("训练数据", "GroundingCap-1M；包含 grounding 标签和详细图像 caption，使用 COCO/LVIS、Flickr30k Entities、GQA、LLaVA-ReCap、V3Det 等底图。"),
            ("训练集", "作者生成的 JSONL 公开；底图需从各原始数据集获取。"),
            ("代码", "完整训练、评测与 Transformers 用法公开。"),
            ("权重", "Swin-T、Swin-B、Swin-L checkpoint 与日志在 Hugging Face / ModelScope 公布。"),
            ("许可", "Apache-2.0；各底层数据集许可证独立。"),
            ("落地难度", "高（4/5 推理，5/5 复现训练）：依赖 MMDetection、SigLIP、轻量 LLaVA/LLM 组件；官方训练示例为 8 GPU。"),
        ],
        "LLMDet 的核心贡献是用详细描述监督弥补传统开放词汇检测只学短类别词的问题。它更能理解属性、动作和场景语义，但最终任务仍偏开放词汇检测，并不天然保证 leftmost、between、second from left 等实例级关系选择。",
        [
            "对 AIC 的长 Query、属性和动作描述具有潜在价值。",
            "可作为强语义候选生成器，与 PIZA/APE 分工；不建议立即作为唯一最终模型。",
            "不使用 IR/Depth，且训练成本远高于当前 8GB 显存环境的舒适区。",
        ],
        [
            "只先测 LLMDet Swin-T（或 only-p5 轻配置），不要从 B/L 开始。",
            "在统一外部验证上比较其 Top-1、Top-K 和对长 Query/attribute/action 的独有召回。",
            "若互补性明显，把其候选加入 union；先不训练 selector，直接做 oracle 与保守 gating。",
            "只有在推理结果证明价值后，再评估冻结大部分模块的 LoRA/adapter 微调。",
        ],
        [
            "复杂依赖和多组件 checkpoint 增加 Windows 复现风险。",
            "GroundingCap-1M 的图像来源多样，外部 benchmark 结果可能含数据重叠。",
            "它的“LLM supervision”是训练方法，不意味着推理时需要商业在线 API；正式方案应只使用本地开源权重。",
        ],
        [
            ("CVPR 论文", "https://openaccess.thecvf.com/content/CVPR2025/html/Fu_LLMDet_Learning_Strong_Open-Vocabulary_Object_Detectors_under_the_Supervision_of_CVPR_2025_paper.html"),
            ("官方代码/权重", "https://github.com/iSEE-Laboratory/LLMDet"),
        ],
        "适合做 P2 语义对照，不是下一轮第一训练主线。先验证 Swin-T 的独有召回，再决定是否承担工程成本。",
    )

    add_paper_section(
        doc,
        "2.5",
        "MM-Grounding-DINO：完整可训练的统一 Grounding/Detection 流水线",
        [
            ("发表状态", "2024 arXiv 论文 + OpenMMLab/MMDetection 官方开源工程"),
            ("任务", "OVD、Phrase Grounding、REC；RGB + text → bbox"),
            ("核心目标", "补齐原始 GroundingDINO 训练细节不足的问题，提供统一的公开训练、微调和评测管线。"),
            ("训练数据", "提供 Objects365、GoldG、GRIT、V3Det 等多种预训练组合及 RefCOCO 等微调配置。"),
            ("训练集", "配置与转换流程公开；各数据集需按官方许可下载。"),
            ("代码", "MMDetection 主仓库中完整公开。"),
            ("权重", "MM-GDINO-T/B/L 多个 checkpoint 与日志公开。"),
            ("许可", "继承 MMDetection 的 Apache-2.0；数据条款独立。"),
            ("落地难度", "中高（3/5 推理，4/5 微调）：Tiny 适合先做；B/L 对 8GB 不友好。"),
        ],
        "MM-GDINO 的最大价值是工程可控：可以明确知道冻结哪些模块、使用哪些训练数据、如何加入新数据集。官方结果显示 Tiny 在多项 OVD 和 RefEXP 指标上优于或接近原 GroundingDINO-Tiny，但这些指标仍不能替代 AIC 平台验证。",
        [
            "最适合作为后续真正微调的 RGB+语言基座，而不是继续维护不可训练的外部封装。",
            "可方便引入 SOREC、gRefCOCO 与 RGBT-GroundBench 的 RGB 分支，建立多域训练。",
            "它不原生支持 IR/Depth；跨模态需要额外视觉分支或候选级晚融合。",
        ],
        [
            "先用公开 MM-GDINO-T checkpoint 重跑统一验证和 AIC 提交，确定是否仅更换基座就有收益。",
            "若 Tiny 有互补召回，先做 2k–5k tracer-bullet 微调：batch 1、梯度累积、冻结视觉 backbone 与文本 encoder。",
            "训练目标同时记录 Top-1 与 Top-K oracle；不能只看 loss 或 RefCOCO 指标。",
            "微调后必须重新生成候选并重新校准 selector，不能复用旧 Ranker 的候选分布。",
        ],
        [
            "某些 checkpoint 训练数据包含 RefCOCO 或相近 grounding 数据，外部验证会偏乐观。",
            "MMCV/MMEngine/CUDA 版本组合在 Windows 上需要独立环境，不能污染现有稳定 baseline。",
            "全模型混合精度曾在本项目中触发文本增强层 dtype 问题；先做 FP32 smoke，再逐层启用 autocast。",
        ],
        [
            ("论文", "https://arxiv.org/abs/2401.02361"),
            ("官方配置/权重", "https://github.com/open-mmlab/mmdetection/tree/main/configs/mm_grounding_dino"),
        ],
        "P0 可训练基座。先零样本对照，再小规模 adapter/局部微调；不建议直接上全量 287k 或 B/L 大模型。",
        verdict_color=GREEN,
        verdict_fill=LIGHT_GREEN,
    )

    add_paper_section(
        doc,
        "2.6",
        "RGBDT500 与 RDTTrack：RGB + Depth + Thermal 三模态跟踪",
        [
            ("发表状态", "NeurIPS 2025 Datasets and Benchmarks Track"),
            ("任务", "RGB + Depth + TIR 视频 → 单目标 tracking bbox；无自然语言 Query"),
            ("核心目标", "用三种对齐视觉模态提升复杂场景下目标跟踪的鲁棒性。"),
            ("数据规模", "500 个同步三模态视频，约 203.7K 组三模态帧；每帧有目标 bbox。"),
            ("训练集", "公开下载；含 train/test、RGB/depth/infrared 和 ground_truth。"),
            ("代码", "公开，包含训练、测试和评测工具。"),
            ("权重", "RDTTrack 预训练模型及底层 OSTrack foundation model 链接公开。"),
            ("许可", "代码 MIT；数据使用条款需在项目页和下载说明中进一步确认。"),
            ("落地难度", "很高（5/5）：任务定义不同，需把 tracking prompt fusion 改造成 text-conditioned grounding。"),
        ],
        "RDTTrack 先融合 Thermal 与 Depth，再以 prompt 的形式注入预训练 RGB tracker，并使用正交投影约束。它证明三模态辅助可以在不破坏 RGB 主干的情况下接入，但它没有语言编码器，也不是“给 Query 直接找目标”。",
        [
            "视觉模态与 AIC 最接近，适合参考 Depth/TIR 的适配器、可靠性门控和预训练方式。",
            "不适合作为现成 AIC 模型：tracking 的第一帧目标提示与自然语言 Query 完全不同。",
            "可用于学习候选框内部的三模态表征，但要额外引入文本编码和 REC 训练。",
        ],
        [
            "第一阶段只做论文/代码结构研究和数据同源审计，不下载后直接训练。",
            "若确认合规且无测试集重叠，可用 RGBDT500 做“候选区域三模态一致性/可靠性”预训练。",
            "把训练出的视觉融合特征接到 MM-GDINO selector 或 late-fusion head，而不是直接把 tracker 当 detector。",
            "Depth 关系仅对 nearest/farthest/front/behind 等 Query 开启，并区分 AIC PNG 毫米深度与 JPG 未知强度域。",
        ],
        [
            "与 AIC 可能存在同源或相邻帧风险，使用前必须做精确哈希和近重复审计，并建议向赛事方确认。",
            "tracking bbox 通常针对视频中的同一实例，缺少多实例语义消歧和语言关系监督。",
            "原训练配置使用多张 3090Ti；8GB 环境只能做小规模、冻结式迁移。",
        ],
        [
            ("NeurIPS 论文", "https://proceedings.neurips.cc/paper_files/paper/2025/hash/b4962fcd5d4410a9f43ef70f528eedd8-Abstract-Datasets_and_Benchmarks_Track.html"),
            ("官方代码/权重", "https://github.com/xuefeng-zhu5/RDTTrack"),
            ("项目页", "https://xuefeng-zhu5.github.io/RGBDT500/"),
        ],
        "重要的长期融合参考，但不应被描述为“可直接用于 AIC 的三模态语义检测器”。在许可和重叠风险澄清前暂不进入训练。",
        verdict_color=AMBER,
        verdict_fill=LIGHT_AMBER,
    )

    add_paper_section(
        doc,
        "2.7",
        "Thermo-VL：文本引导的 RGB-Thermal VLM",
        [
            ("发表状态", "2026 arXiv 预印本；本次未核验到正式顶会录用信息"),
            ("任务", "RGB + Thermal + prompt → 文本回答/VQA；不是 bbox grounding"),
            ("核心目标", "在冻结 Molmo-7B RGB 视觉语言接口的前提下，引入可训练 thermal encoder 和文本引导双注意力融合。"),
            ("训练数据", "论文称构建对齐 RGB-Thermal instruction 数据和 Thermo-VL-Bench。"),
            ("训练集", "论文/项目页声称公开，但当前项目页的 Code/Data 未解析到可点击仓库；需后续再次核验。"),
            ("代码", "未确认可访问的官方仓库链接。"),
            ("权重", "未确认公开 checkpoint。"),
            ("许可", "未确认。"),
            ("落地难度", "极高（5/5）：7B VLM、无 bbox head、需要新增 grounding 头和监督。"),
        ],
        "Thermo-VL 的结构思想值得借鉴：thermal tokens 同时关注 Query 与 RGB tokens，再预测 gated residual 注入冻结 RGB 流。这是一种“由语言决定何时使用红外”的动态融合方式，正好对应 AIC 需求。",
        [
            "概念相关：文本引导的模态可靠性融合比固定拼接更合理。",
            "任务不匹配：其输出是语言回答，不是归一化 bbox。",
            "当前 8GB 显存和开放资产状态都不支持把它作为下一轮主线。",
        ],
        [
            "只提取其 dual-attention + gated residual 设计，作为未来 RGBT-VGNet/MM-GDINO 融合模块的参考。",
            "若未来代码与权重真正公开，再在小规模 RGB-TIR validation 上验证，不直接接 AIC 全量。",
            "任何训练数据生成不得复刻其商业闭源模型标注流程；比赛方案只使用合规公开数据和本地开源模型。",
        ],
        [
            "项目页明确提到部分 QA 使用 GPT-5-mini/GPT-5.1 生成；复刻该流程可能与比赛禁止商业闭源 API 的边界冲突。",
            "没有 bbox head 与 REC 训练，改造工作量远超 PIZA、APE 或 MM-GDINO。",
            "不要把“论文说公开”写成“已验证有可下载权重”。",
        ],
        [
            ("arXiv", "https://arxiv.org/abs/2605.21882"),
            ("项目页", "https://rusiru.us/Thermo-VL/"),
        ],
        "仅作为架构参考，暂不进入模型下载、训练或平台提交队列。",
        verdict_color=RED,
        verdict_fill=LIGHT_RED,
    )

    doc.add_page_break()
    doc.add_heading("3. 其他开源候选模型对照", level=1)
    add_body(doc, "以下模型多数是开放词汇目标检测器，而不是完整 Referring Expression Comprehension 模型。它们的合理角色是提供互补候选、速度基线或类别召回，复杂空间/序数关系仍需由专门的关系模块、VLM 或经过验证的 selector 处理。")

    add_table(
        doc,
        ["模型", "发表/状态", "数据/代码/权重", "AIC 角色", "难度与建议"],
        [
            ("OV-DINO", "2024 arXiv；官方实现", "Apache-2.0；O365/GoldG/CC1M 组合权重、推理与微调代码公开", "更强 OVD 候选源；Language-Aware Selective Fusion 可参考", "中高。老版 CUDA/Detectron2 栈；先离线小集验证"),
            ("YOLO-World V2.1", "CVPR 2024", "GPL-3.0；S/M/L/X 多分辨率权重、预训练/微调/部署公开", "高速候选、小目标 1280 分辨率对照", "低中。优先 S-640/S-1280；不承担关系推理"),
            ("OmDet-Turbo-Tiny", "2024 arXiv 工程", "Apache-2.0；Tiny 权重公开并集成 Transformers", "轻量快速候选与速度基线", "低。适合 RTX 4060；只有 Tiny 官方权重"),
            ("OWLv2 Base", "2023 论文/Google 模型", "Apache-2.0；HF 权重与 Transformers 接口公开，约 0.2B 参数", "独立文本条件候选源、长尾补充", "低中。易集成；完整 Query 关系选择有限"),
            ("GroundingDINO-Tiny/B", "ECCV 2024", "官方/HF 权重公开；Tiny 已在本项目验证", "现有候选源；B 可做受控大模型对照", "Tiny 低；B 中高。平台 Top-1 0.4938"),
            ("Florence-2-large-ft", "2024 技术报告/模型卡", "MIT；模型权重公开，当前工程已集成", "稳定 RGB baseline、候选/生成式对照", "低。平台 0.4980，目前仍是最佳实测基线"),
        ],
        widths=[1.15, 1.1, 1.8, 1.25, 1.2],
        font_size=7.8,
    )

    doc.add_heading("3.1 OV-DINO", level=2)
    add_body(doc, "OV-DINO 通过 Unified Data Integration 与 Language-Aware Selective Fusion 提升开放词汇检测。官方仓库提供 Objects365、Objects365+GoldG、再加 CC1M 的多组 Swin-T 权重，也提供自定义 COCO 格式微调。它比现有 GDINO 更适合测试“换候选生成器能否提供独有召回”，但官方环境依赖 PyTorch 1.13.1、CUDA 11.6 和 Detectron2/detrex，Windows 复现成本不低。")
    add_bullet(doc, "建议：只做 500–1,500 条统一验证候选缓存，先比较 union oracle；未证明互补前不跑 9,555 全量。")
    add_link_line(doc, [("官方仓库", "https://github.com/wanghao9610/OV-DINO"), ("模型仓库", "https://huggingface.co/hao9610/OV-DINO"), ("论文", "https://arxiv.org/abs/2407.07844")])

    doc.add_heading("3.2 YOLO-World", level=2)
    add_body(doc, "YOLO-World 的优势是速度、部署和高分辨率模型。官方 V2.1 提供 640/1280 分辨率权重，1280 版本可作为小目标候选对照。它本质上更接近 prompt-then-detect：给类别/短语后快速检测，不擅长直接理解“左数第二个、位于公交车左侧的人”。")
    add_bullet(doc, "建议：用 target phrase 而不是整句 Query 产生候选，再交给独立关系模块；注意 GPL-3.0 对分发方式的影响。")
    add_link_line(doc, [("CVPR 论文", "https://openaccess.thecvf.com/content/CVPR2024/papers/Cheng_YOLO-World_Real-Time_Open-Vocabulary_Object_Detection_CVPR_2024_paper.pdf"), ("官方仓库", "https://github.com/AILab-CVC/YOLO-World")])

    doc.add_heading("3.3 OmDet-Turbo", level=2)
    add_body(doc, "OmDet-Turbo 使用 Efficient Fusion Head 降低开放词汇检测的融合开销。官方只公开 OmDet-Turbo-Tiny 权重，但已经集成 Transformers，部署门槛低。论文层面属于 2024 arXiv 预印本，不能与已正式录用的 CVPR/ICCV 模型同级表述。")
    add_bullet(doc, "建议：作为最轻量的 candidate-union 补充分支或速度下限，不作为复杂语义主模型。")
    add_link_line(doc, [("官方仓库", "https://github.com/om-ai-lab/OmDet"), ("HF 模型", "https://huggingface.co/omlab/omdet-turbo-swin-tiny-hf"), ("论文", "https://arxiv.org/abs/2403.06892")])

    doc.add_heading("3.4 OWLv2", level=2)
    add_body(doc, "OWLv2 是 Google 的文本条件零样本检测模型，HF 的 base-patch16-ensemble 约 0.2B 参数，Apache-2.0，使用标准 Transformers 接口。它通过大规模自训练提高开放词汇检测，适合提供与 GDINO 不同的候选分布。")
    add_bullet(doc, "建议：用目标实体短语跑 Top-K，评估独有救回率；不要期待它自动解决 reference relation 或序数。")
    add_link_line(doc, [("HF 模型", "https://huggingface.co/google/owlv2-base-patch16-ensemble"), ("论文", "https://arxiv.org/abs/2306.09683")])

    doc.add_heading("3.5 现有 Florence-2 与 GroundingDINO 的正确定位", level=2)
    add_body(doc, "平台结果说明二者是同一量级：Florence 0.4980，GroundingDINO-Tiny 0.4938。Florence 仍是稳定控制组；GroundingDINO 仍可作为候选源，但不能再附加当前无保护的 Spatial LTR。后续任何新模型都必须与这两个固定提交做单变量比较。")
    add_callout(doc, "当前 Ranker 状态", "RefCOCO Spatial LTR 已被 AIC 平台证伪：切换近一半预测、偏向大框、频繁从目标翻到参照物。可以保留代码和诊断价值，但不得作为下一轮默认 selector。", color=RED, fill=LIGHT_RED)

    doc.add_page_break()
    doc.add_heading("4. 面向 AIC 的组合方案与训练路线", level=1)
    doc.add_heading("4.1 推荐系统不是一个万能模型，而是受控分工", level=2)
    add_table(
        doc,
        ["分支", "输入", "主要负责", "优先实现"],
        [
            ("稳定主分支", "RGB + Query", "通用视觉定位", "Florence-2 / MM-GDINO-T"),
            ("小目标分支", "RGB + Query", "无人机、摄像头、灯泡、远处标志", "PIZA"),
            ("区域结构分支", "RGB + Query", "passage、awning、entrance、stuff/region", "APE-Ti"),
            ("跨光谱分支", "RGB + IR + Query", "弱光、热目标、恶劣天气", "RGBT-VGNet"),
            ("快速候选分支", "RGB + target phrase", "补充长尾类别和漏检候选", "OmDet-Turbo / YOLO-World / OWLv2"),
            ("Depth 决策层", "Depth + candidate boxes + Query flags", "nearest/farthest/front/behind", "候选级 late fusion，最后加入"),
        ],
        widths=[1.1, 1.35, 2.4, 1.65],
        font_size=8.4,
    )
    add_body(doc, "推荐的信息流是：各分支只产生候选和自身置信度；统一模块先做目标/参照物/部件角色过滤，再进行保守选择。若新分支不能以足够置信度证明自己更好，则回退到 Florence 或原模型 Top-1。")

    doc.add_heading("4.2 新 selector 的训练原则", level=2)
    add_bullet(doc, "训练域：RefCOCO 系列只占一部分；必须加入 SOREC（极小目标）、gRefCOCO（复数/群组）、RGBT-GroundBench（低光/热红外）和区域类数据。")
    add_bullet(doc, "特征：candidate tight crop、context crop、全图语义、目标短语/参照物/部件角色，而不是依赖绝对 area/log_area。")
    add_bullet(doc, "目标：优先预测 P(IoU ≥ 0.5)，并把参照物框、部件框、包含大区域和同标签错误实例作为 hard negatives。")
    add_bullet(doc, "保护：限制跨 label 切换、GroundingDINO score 降幅、面积倍率和目标到参照物翻转；低置信 Query 不切换。")
    add_bullet(doc, "验证：必须按数据域、Query 类型和面积区间报告；不能再只看 RefCOCO 总分。")

    doc.add_heading("4.3 建议的数据组合", level=2)
    add_table(
        doc,
        ["数据", "主要监督", "可用性", "在 AIC 中的用途", "注意事项"],
        [
            ("RefCOCO / + / g", "普通实体、属性、关系", "已有", "保持通用 REC 基础", "不可主导训练分布"),
            ("SOREC", "极小目标 + 长关系句", "标注/adapter 公开", "小目标分支训练与验证", "底图需 SODA-D"),
            ("gRefCOCO", "复数/群组与无目标", "已有/可准备", "组合框与群组监督", "标注定义需统一"),
            ("RGBT-GroundBench", "RGB-TIR + Query + bbox", "公开", "跨光谱训练与低光验证", "许可、同源/重叠审计"),
            ("RGBDT500", "RGB-D-T tracking bbox", "公开", "融合预训练/候选表征", "无 Query；暂缓训练"),
            ("AIC 9,555", "无 bbox", "仅测试", "只做推理和提交", "禁止训练、人工标注和伪标签反训"),
        ],
        widths=[1.25, 1.25, 1.05, 1.45, 1.5],
        font_size=8.2,
    )

    doc.add_heading("4.4 RTX 4060 约束下的模型选择", level=2)
    add_bullet(doc, "优先：PIZA adapter、APE-Ti、MM-GDINO-T、OmDet-Turbo-Tiny、YOLO-World-S、OWLv2 Base。")
    add_bullet(doc, "谨慎：LLMDet Swin-T、OV-DINO、RGBT-VGNet；建议独立环境、先 smoke、缓存候选。")
    add_bullet(doc, "暂不做：APE-L、MM-GDINO-B/L 全量微调、LLMDet B/L、Thermo-VL 7B 端到端训练。")
    add_bullet(doc, "训练统一采用：batch size 1、梯度累积、gradient checkpointing、冻结 backbone/text encoder、先 FP32 smoke，再选择性混合精度。")

    doc.add_page_break()
    doc.add_heading("5. 推荐实验顺序、验收标准与风险控制", level=1)
    add_table(
        doc,
        ["阶段", "动作", "产物", "晋级条件"],
        [
            ("E0 资产审计", "固定官方 repo/commit、checkpoint、配置、许可证与 SHA-256", "sota_asset_manifest.json / md", "所有权重与配置一一对应；许可/来源无歧义"),
            ("E1 统一外部验证", "Florence、GDINO、PIZA、APE-Ti、MM-GDINO-T、OmDet/YOLO/OWL", "多域 ACC@0.5、Top-K、速度、显存", "至少一个模型在目标难点上产生稳定独有召回"),
            ("E2 单模型平台", "只提交最强新单模型 Top-1", "一个 ZIP + 审计 + 哈希", "超过 0.4980，或明确证明某子任务价值"),
            ("E3 候选互补", "计算 union oracle 与重复率；先不训练 selector", "candidate complement report", "union 提升明显且候选规模可控"),
            ("E4 保守 selector", "目标角色、视觉语义、hard negatives、切换保护", "受控 Ranker v2", "多域验证不出现大域崩溃；harm 明显受控"),
            ("E5 RGB+IR", "复现/适配 RGBT-VGNet", "RGB/TIR/RGBT 消融", "官方 val 可复现；AIC 单变量提交优于 RGB 控制"),
            ("E6 Depth", "只对明确深度词做 candidate late fusion", "Depth 单变量提交", "PNG/JPG 域均有独立审计，不伪造二维关系"),
        ],
        widths=[0.95, 2.25, 1.65, 1.65],
        font_size=8.1,
    )

    doc.add_heading("下一次最合理的三项实际动作", level=2)
    add_numbered(doc, 1, "建立统一模型资产清单，并优先把 PIZA adapter、APE-Ti、MM-GDINO-T、OmDet-Turbo-Tiny 下载到独立缓存；不污染现有 baseline 环境。")
    add_numbered(doc, 2, "用同一批 RefCOCO/SOREC/gRefCOCO/RGBT-GroundBench validation manifest 跑模型对照，报告 Top-1、Top-K、按面积和 Query 类型分组结果。")
    add_numbered(doc, 3, "只从结果中选择一个变量生成下一份 AIC 提交：优先“PIZA 小目标门控”或“最佳新单模型 Top-1”，不要再次一次性叠加 Ranker、Tile、IR、Depth。")

    doc.add_heading("明确暂缓的工作", level=2)
    add_bullet(doc, "继续扩大当前 LightGBM Spatial LTR 训练规模。")
    add_bullet(doc, "在没有多域验证前全量微调 GroundingDINO 或直接更换到大型 B/L backbone。")
    add_bullet(doc, "把 RGB、IR、Depth 简单拼成 5 通道并端到端训练。")
    add_bullet(doc, "使用 AIC 正式测试图像、人工 bbox 或测试伪标签进行训练。")
    add_bullet(doc, "在没有许可证与近重复审计前使用 RGBDT500 或与赛题可能同源的数据。")

    add_callout(doc, "最重要的判断标准", "一个模型是否值得进入 AIC，不看它在自己的论文表格里是不是 SOTA，而看它是否在固定多域验证上补回 AIC 的真实难点，并能通过一次单变量平台提交证明迁移。", color=BLUE, fill=LIGHT_BLUE)

    doc.add_page_break()
    doc.add_heading("6. 官方来源与链接索引", level=1)
    sources = [
        ("[L1] 本项目 Spatial LTR 平台复盘", str(ROOT / "reports" / "gdino_spatial_ltr_v1_platform_postmortem.md")),
        ("[S1] RGBT-GroundBench arXiv", "https://arxiv.org/abs/2512.24561"),
        ("[S2] RGBT-GroundBench 官方仓库", "https://github.com/crazyxiaoxi/RGBT-GroundBench"),
        ("[S3] RGBT-GroundBench 数据", "https://huggingface.co/datasets/JiawenXi/RGBT-Ground-Dataset"),
        ("[S4] RGBT-Ground-Model", "https://huggingface.co/JiawenXi/RGBT-Ground-Model"),
        ("[S5] PIZA/SOREC 官方仓库", "https://github.com/mmaiLab/sorec"),
        ("[S6] PIZA/SOREC ICCV 2025 论文", "https://openaccess.thecvf.com/content/ICCV2025/papers/Goto_Referring_Expression_Comprehension_for_Small_Objects_ICCV_2025_paper.pdf"),
        ("[S7] APE CVPR 2024 论文", "https://openaccess.thecvf.com/content/CVPR2024/html/Shen_Aligning_and_Prompting_Everything_All_at_Once_for_Universal_Visual_CVPR_2024_paper.html"),
        ("[S8] APE 官方仓库与权重", "https://github.com/shenyunhang/APE"),
        ("[S9] LLMDet CVPR 2025 论文", "https://openaccess.thecvf.com/content/CVPR2025/html/Fu_LLMDet_Learning_Strong_Open-Vocabulary_Object_Detectors_under_the_Supervision_of_CVPR_2025_paper.html"),
        ("[S10] LLMDet 官方仓库与权重", "https://github.com/iSEE-Laboratory/LLMDet"),
        ("[S11] MM-Grounding-DINO 论文", "https://arxiv.org/abs/2401.02361"),
        ("[S12] MM-Grounding-DINO 官方配置与权重", "https://github.com/open-mmlab/mmdetection/tree/main/configs/mm_grounding_dino"),
        ("[S13] RGBDT500/RDTTrack NeurIPS 2025", "https://proceedings.neurips.cc/paper_files/paper/2025/hash/b4962fcd5d4410a9f43ef70f528eedd8-Abstract-Datasets_and_Benchmarks_Track.html"),
        ("[S14] RDTTrack 官方仓库", "https://github.com/xuefeng-zhu5/RDTTrack"),
        ("[S15] Thermo-VL arXiv", "https://arxiv.org/abs/2605.21882"),
        ("[S16] Thermo-VL 项目页", "https://rusiru.us/Thermo-VL/"),
        ("[S17] OV-DINO 官方仓库", "https://github.com/wanghao9610/OV-DINO"),
        ("[S18] YOLO-World 官方仓库", "https://github.com/AILab-CVC/YOLO-World"),
        ("[S19] OmDet-Turbo 官方仓库", "https://github.com/om-ai-lab/OmDet"),
        ("[S20] OWLv2 Hugging Face 模型卡", "https://huggingface.co/google/owlv2-base-patch16-ensemble"),
    ]
    for label, url in sources:
        p = doc.add_paragraph(style="Small Note")
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(label + " — ")
        r.bold = True
        set_east_asia_font(r)
        if url.startswith("http"):
            add_hyperlink(p, url, url)
        else:
            r2 = p.add_run(url)
            r2.font.name = "Consolas"
            r2.font.size = Pt(8.5)
            r2._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    doc.add_heading("附注：如何解读“公开”", level=2)
    add_bullet(doc, "代码公开不等于训练可立即复现：仍需要对应配置、底层数据、checkpoint 和依赖版本。")
    add_bullet(doc, "权重仓库存在不等于模型卡完整：必须核对 checkpoint 文件、网络配置、预处理与输出格式。")
    add_bullet(doc, "论文使用公开数据不等于数据可任意再分发或商用：每个底层数据集许可仍然有效。")
    add_bullet(doc, "论文中的 benchmark 分数不等于 AIC 可获得的分数；最终结论必须以 AIC 单变量平台提交为准。")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("— 完 —")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)
    set_east_asia_font(r)
    return doc


def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = build_document()
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
